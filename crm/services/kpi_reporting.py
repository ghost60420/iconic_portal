import csv
from dataclasses import asdict, dataclass
from io import BytesIO, StringIO
from xml.sax.saxutils import escape

from django.utils import timezone
from django.utils.html import strip_tags
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from crm.services.kpi_dashboard import AUDIENCE_EMPLOYEE, dashboard_audience
from crm.services.kpi_intelligence import (
    INTELLIGENCE_VERSION,
    REPORT_TYPES,
    KPIIntelligencePermissionError,
    allowed_report_types,
    intelligence_widget_payload,
)


REPORT_SOURCE_NOTE = (
    "Source: digest-verified approved KPI review snapshots and immutable "
    "Stage 6 bonus results. Historical scores were not recalculated."
)


@dataclass(frozen=True, slots=True)
class KPIReportDocument:
    report_type: str
    title: str
    date_range: str
    filters: tuple[tuple[str, str], ...]
    generated_at: str
    generated_by: str
    scope: str
    summary_score: str
    summary_status: str
    columns: tuple[str, ...]
    rows: tuple[tuple[object, ...], ...]
    main_risks: tuple[str, ...]
    recommended_actions: tuple[str, ...]
    data_source_note: str = REPORT_SOURCE_NOTE
    intelligence_version: str = INTELLIGENCE_VERSION

    def as_dict(self):
        return asdict(self)


def _clean(value, limit=1000):
    text = " ".join(strip_tags(str(value or "")).split())
    return text[:limit]


def _spreadsheet_value(value, limit=1000):
    text = _clean(value, limit)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _filter_rows(filters):
    labels = {
        "employee_id": "Employee ID",
        "role_id": "KPI Role ID",
        "department": "Department",
        "location": "Location",
        "period_type": "Review Period",
        "month": "Month",
        "quarter": "Quarter",
        "year": "Year",
        "manager_id": "Manager ID",
        "status": "Status",
    }
    return tuple(
        (labels[key], _clean(value))
        for key, value in asdict(filters).items()
        if value not in (None, "")
    )


def _date_range(filters):
    if filters.month:
        return f"{filters.year}-{filters.month:02d}"
    if filters.quarter:
        return f"Q{filters.quarter} {filters.year}"
    return str(filters.year)


def _user_name(user):
    profile = getattr(user, "employee_profile", None)
    if profile and profile.public_name:
        return profile.public_name
    return user.get_full_name() or user.get_username()


def _report_scope(user, filters):
    audience = dashboard_audience(user)
    if audience == AUDIENCE_EMPLOYEE:
        return "Own approved KPI records"
    if filters.department:
        return f"Department: {_clean(filters.department, 80)}"
    return f"{audience.title()} authorized scope"


def _health_summary(user, filters):
    try:
        payload = intelligence_widget_payload(
            user,
            "company-health",
            filters,
        )
    except KPIIntelligencePermissionError:
        payload = intelligence_widget_payload(
            user,
            "employee-analytics",
            filters,
        )
        first = payload.get("rows", [{}])[0] if payload.get("rows") else {}
        return first.get("current_score", ""), first.get("status", "")
    return payload.get("score", ""), payload.get("status", "")


def _risks_and_actions(user, filters):
    try:
        risks = intelligence_widget_payload(
            user,
            "critical-alerts",
            filters,
        ).get("rows", [])
        actions = intelligence_widget_payload(
            user,
            "recommended-actions",
            filters,
        ).get("rows", [])
    except KPIIntelligencePermissionError:
        risks = intelligence_widget_payload(
            user,
            "yellow-attention",
            filters,
        ).get("rows", [])
        actions = risks
    return (
        tuple(_clean(row.get("summary"), 300) for row in risks[:5]),
        tuple(
            _clean(row.get("recommended_action"), 300)
            for row in actions[:5]
            if row.get("recommended_action")
        ),
    )


def _table_for_report(user, report_type, filters):
    if report_type == "executive-summary":
        payload = intelligence_widget_payload(user, "company-health", filters)
        columns = (
            "Company Health",
            "Status",
            "Overall KPI",
            "Review Completion",
            "Green",
            "Yellow",
            "Red",
            "Critical Red",
            "Bonus Readiness",
        )
        rows = (
            (
                payload.get("score"),
                payload.get("status"),
                payload.get("overall_kpi"),
                payload.get("review_completion"),
                payload.get("green"),
                payload.get("yellow"),
                payload.get("red"),
                payload.get("critical_red"),
                payload.get("bonus_readiness"),
            ),
        )
    elif report_type == "department-performance":
        payload = intelligence_widget_payload(
            user,
            "department-analytics",
            filters,
        )
        columns = (
            "Department",
            "Current Score",
            "Previous Score",
            "Direction",
            "Green",
            "Yellow",
            "Red",
            "Review Completion",
            "Top Strength",
            "Main Risk",
            "Recommended Action",
        )
        rows = tuple(
            (
                row.get("name"),
                row.get("current_score"),
                row.get("previous_score"),
                row.get("trend", {}).get("direction"),
                row.get("green"),
                row.get("yellow"),
                row.get("red"),
                row.get("review_completion"),
                row.get("top_strength"),
                row.get("main_risk"),
                row.get("recommended_action"),
            )
            for row in payload.get("rows", [])
        )
    elif report_type == "manager-performance":
        payload = intelligence_widget_payload(
            user,
            "manager-analytics",
            filters,
        )
        columns = (
            "Manager",
            "Team Average",
            "Review Completion",
            "Overdue Reviews",
            "Green",
            "Yellow",
            "Red",
            "Team Improvement",
            "Missed KPI Items",
            "Bonus Ready",
            "Workload",
        )
        rows = tuple(
            (
                row.get("manager"),
                row.get("team_average"),
                row.get("review_completion"),
                row.get("overdue_reviews"),
                row.get("green"),
                row.get("yellow"),
                row.get("red"),
                row.get("team_improvement"),
                row.get("repeated_missed_items"),
                row.get("bonus_ready"),
                row.get("workload"),
            )
            for row in payload.get("rows", [])
        )
    elif report_type == "employee-performance":
        payload = intelligence_widget_payload(
            user,
            "employee-analytics",
            filters,
            page=None,
            use_cache=False,
        )
        columns = (
            "Employee",
            "Current Score",
            "Previous Score",
            "Change",
            "Status",
            "Strongest KPI",
            "Needs Improvement",
            "Review Status",
            "Review Date",
            "Approved Comments",
            "Suggested Action",
        )
        rows = tuple(
            (
                row.get("employee"),
                row.get("current_score"),
                row.get("previous_score"),
                row.get("change"),
                row.get("status"),
                row.get("strongest_kpi"),
                row.get("improvement_kpi"),
                row.get("review_status"),
                row.get("review_date"),
                " | ".join(row.get("approved_comments", ())),
                row.get("suggested_action"),
            )
            for row in payload.get("rows", [])
        )
    elif report_type == "review-completion":
        payload = intelligence_widget_payload(
            user,
            "review-completion",
            filters,
        )
        columns = (
            "Completion",
            "Completed",
            "Missing",
            "Open",
            "Overdue",
            "Target",
            "Target Met",
        )
        rows = (
            (
                payload.get("completion"),
                payload.get("completed"),
                payload.get("missing"),
                payload.get("open"),
                payload.get("overdue"),
                payload.get("target"),
                "Yes" if payload.get("target_met") else "No",
            ),
        )
    elif report_type == "critical-red":
        payload = intelligence_widget_payload(
            user,
            "critical-alerts",
            filters,
            page=None,
            use_cache=False,
        )
        columns = (
            "Title",
            "Reason",
            "Severity",
            "Employee",
            "Manager",
            "Department",
            "Review Period",
            "Supporting Value",
            "Recommended Action",
        )
        rows = tuple(
            (
                row.get("title"),
                row.get("summary"),
                row.get("severity"),
                row.get("employee"),
                row.get("manager"),
                row.get("department"),
                row.get("review_period"),
                row.get("supporting_value"),
                row.get("recommended_action"),
            )
            for row in payload.get("rows", [])
        )
    elif report_type == "bonus-readiness":
        payload = intelligence_widget_payload(
            user,
            "bonus-readiness",
            filters,
            page=None,
            use_cache=False,
        )
        columns = (
            "Employee",
            "Status",
            "Reason",
            "Critical Red",
            "Rule Version",
            "Calculation Date",
            "Authorized Amount",
        )
        rows = tuple(
            (
                row.get("employee"),
                row.get("status"),
                row.get("reason"),
                "Yes" if row.get("critical_red") else "No",
                row.get("rule_version"),
                row.get("calculation_date"),
                row.get("amount"),
            )
            for row in payload.get("rows", [])
        )
    else:
        period = report_type.split("-", 1)[0]
        payload = intelligence_widget_payload(user, "trends", filters)
        trend = payload.get(period, {})
        columns = (
            "Period",
            "Approved KPI Average",
            "Current",
            "Previous",
            "Difference",
            "Percentage Change",
            "Direction",
            "Status",
            "History Warning",
        )
        analysis = trend.get("analysis", {})
        rows = tuple(
            (
                point.get("label"),
                point.get("value"),
                analysis.get("current_result"),
                analysis.get("previous_result"),
                analysis.get("difference"),
                analysis.get("percentage_change"),
                analysis.get("direction"),
                analysis.get("status"),
                analysis.get("data_quality_warning"),
            )
            for point in trend.get("points", [])
        )
    return tuple(columns), tuple(tuple(_clean(value) for value in row) for row in rows)


def build_kpi_report(user, report_type, filters):
    allowed = {row["code"] for row in allowed_report_types(user)}
    if report_type not in allowed or report_type not in REPORT_TYPES:
        raise KPIIntelligencePermissionError("This KPI report is not available.")
    columns, rows = _table_for_report(user, report_type, filters)
    summary_score, summary_status = _health_summary(user, filters)
    risks, actions = _risks_and_actions(user, filters)
    return KPIReportDocument(
        report_type=report_type,
        title=REPORT_TYPES[report_type],
        date_range=_date_range(filters),
        filters=_filter_rows(filters),
        generated_at=timezone.now().isoformat(),
        generated_by=_clean(_user_name(user), 160),
        scope=_report_scope(user, filters),
        summary_score=_clean(summary_score, 40),
        summary_status=_clean(summary_status, 40),
        columns=columns,
        rows=rows,
        main_risks=risks,
        recommended_actions=actions,
    )


def report_csv_bytes(report):
    output = StringIO(newline="")
    writer = csv.writer(output)
    safe_write = lambda row: writer.writerow(
        [_spreadsheet_value(value) for value in row]
    )
    _write_metadata_rows(safe_write, report)
    writer.writerow(())
    safe_write(report.columns)
    for row in report.rows:
        safe_write(row)
    return output.getvalue().encode("utf-8-sig")


def report_xlsx_bytes(report):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "KPI Report"
    gold = "D6B45A"
    dark = "15181D"
    sheet.append([report.title])
    sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=dark)
    metadata = (
        ("Date range", report.date_range),
        ("Generated", report.generated_at),
        ("Generated by", report.generated_by),
        ("Scope", report.scope),
        ("Summary score", report.summary_score),
        ("Status", report.summary_status),
        ("Data source", report.data_source_note),
    )
    for label, value in metadata:
        sheet.append([_spreadsheet_value(label), _spreadsheet_value(value)])
    if report.filters:
        sheet.append(
            [
                "Filters",
                _spreadsheet_value(
                    "; ".join(
                        f"{key}: {value}" for key, value in report.filters
                    )
                ),
            ]
        )
    sheet.append([])
    sheet.append(list(report.columns))
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="111111")
        cell.fill = PatternFill("solid", fgColor=gold)
    for row in report.rows:
        sheet.append([_spreadsheet_value(value) for value in row])
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = (
        f"A{header_row}:{sheet.cell(header_row + len(report.rows), len(report.columns)).coordinate}"
        if report.columns and report.rows
        else f"A{header_row}:{sheet.cell(header_row, max(len(report.columns), 1)).coordinate}"
    )
    for column_cells in sheet.columns:
        width = min(
            max(len(str(cell.value or "")) for cell in column_cells) + 2,
            50,
        )
        sheet.column_dimensions[column_cells[0].column_letter].width = max(width, 12)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def report_pdf_bytes(report):
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        rightMargin=0.35 * inch,
        leftMargin=0.35 * inch,
        topMargin=0.35 * inch,
        bottomMargin=0.35 * inch,
        title=report.title,
        author=report.generated_by,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(escape(_clean(report.title)), styles["Title"]),
        Paragraph(
            escape(
                _clean(
                    f"{report.date_range} | Generated {report.generated_at} | "
                    f"By {report.generated_by}"
                )
            ),
            styles["BodyText"],
        ),
        Paragraph(
            escape(
                _clean(
                    f"Scope: {report.scope} | Summary: "
                    f"{report.summary_score or 'Unavailable'} "
                    f"{report.summary_status.title() if report.summary_status else ''}"
                )
            ),
            styles["BodyText"],
        ),
        Spacer(1, 0.14 * inch),
    ]
    table_rows = [
        [
            Paragraph(escape(_clean(column, 120)), styles["BodyText"])
            for column in report.columns
        ]
    ]
    table_rows.extend(
        [
            [
                Paragraph(escape(_clean(value, 300)), styles["BodyText"])
                for value in row
            ]
            for row in report.rows
        ]
    )
    if not report.rows:
        table_rows.append(
            [Paragraph("No authorized results.", styles["BodyText"])]
            + [""] * max(len(report.columns) - 1, 0)
        )
    column_count = max(len(report.columns), 1)
    table = Table(
        table_rows,
        repeatRows=1,
        colWidths=[10.2 * inch / column_count] * column_count,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D6B45A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#111111")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D1D5DB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F8")]),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend(
        [
            table,
            Spacer(1, 0.15 * inch),
            Paragraph(
                "<b>Main risks:</b> "
                + escape(
                    _clean(
                        " | ".join(report.main_risks) or "None identified."
                    )
                ),
                styles["BodyText"],
            ),
            Paragraph(
                "<b>Recommended actions:</b> "
                + escape(
                    _clean(
                        " | ".join(report.recommended_actions)
                        or "Continue monitoring."
                    )
                ),
                styles["BodyText"],
            ),
            Paragraph(
                escape(_clean(report.data_source_note)),
                styles["BodyText"],
            ),
        ]
    )
    document.build(story)
    return output.getvalue()


def _write_metadata_rows(write_row, report):
    write_row((report.title,))
    write_row(("Date range", report.date_range))
    write_row(("Generated", report.generated_at))
    write_row(("Generated by", report.generated_by))
    write_row(("Scope", report.scope))
    write_row(("Summary score", report.summary_score))
    write_row(("Status", report.summary_status))
    if report.filters:
        write_row(
            (
                "Filters",
                "; ".join(f"{key}: {value}" for key, value in report.filters),
            )
        )
    write_row(("Main risks", " | ".join(report.main_risks)))
    write_row(("Recommended actions", " | ".join(report.recommended_actions)))
    write_row(("Data source", report.data_source_note))
