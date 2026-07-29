from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import Group
from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.db import connection
from django.test import override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from crm.models import CRMAuditLog
from crm.models_kpi_bonus import KPIBonusRuleSet, KPIBonusWeightProfile
from crm.models_kpi_intelligence import KPIIntelligenceRuleSet
from crm.models_kpi_reviews import KPIReview
from crm.services.kpi_assignments import assign_kpi_role
from crm.services.kpi_bonus import create_bonus_calculation
from crm.services.kpi_dashboard import DashboardFilters, dashboard_audience
from crm.services.kpi_intelligence import (
    KPIIntelligencePermissionError,
    allowed_report_types,
    effective_intelligence_rule_set,
    intelligence_widget_payload,
    intelligence_widgets,
    publish_intelligence_rule_set,
    resolve_intelligence_action,
)
from crm.services.kpi_reporting import (
    build_kpi_report,
    report_csv_bytes,
    report_pdf_bytes,
    report_xlsx_bytes,
)
from crm.services.kpi_reviews import (
    approve_kpi_review,
    create_kpi_review,
    save_kpi_review_draft,
    start_kpi_review,
    submit_kpi_review,
)
from crm.tests.test_kpi_performance_ui import KPIReviewTestBase


@override_settings(
    CACHES={
        "default": {
            "BACKEND": "django.core.cache.backends.locmem.LocMemCache",
            "LOCATION": "stage8-intelligence-tests",
        }
    }
)
class KPIIntelligenceTests(KPIReviewTestBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.other_assignment = assign_kpi_role(
            employee=cls.other_employee,
            kpi_template=cls.template,
            role_weight=Decimal("100.00"),
            manager=cls.other_manager,
            start_date=date(2026, 1, 1),
            actor=cls.admin,
            reason="Stage 8 production assignment",
        )
        cls.employee_user.groups.add(
            Group.objects.get_or_create(name="CA_TEAM")[0]
        )
        cls.other_employee_user.groups.add(
            Group.objects.get_or_create(name="BD_TEAM")[0]
        )
        cls.intelligence_rules = KPIIntelligenceRuleSet.objects.create(
            code="iconic-intelligence",
            name="Iconic Intelligence Test Rules",
            version=1,
            effective_start=date(2026, 1, 1),
            minimum_score=Decimal("70.00"),
            health_green_threshold=Decimal("85.00"),
            critical_red_count=2,
            decline_percentage=Decimal("5.00"),
            trend_periods=3,
            overdue_days=7,
            review_completion_target=Decimal("90.00"),
            improvement_threshold=Decimal("5.00"),
            department_risk_threshold=Decimal("70.00"),
            manager_workload_threshold=5,
            overall_kpi_weight=Decimal("40.00"),
            review_completion_weight=Decimal("20.00"),
            bonus_readiness_weight=Decimal("15.00"),
            improvement_weight=Decimal("10.00"),
            risk_control_weight=Decimal("15.00"),
            created_by=cls.admin,
        )
        cls.intelligence_rules.status = KPIIntelligenceRuleSet.STATUS_PUBLISHED
        cls.intelligence_rules.published_by = cls.admin
        cls.intelligence_rules.save()

    def setUp(self):
        cache.clear()

    def approved_period(
        self,
        *,
        employee=None,
        manager=None,
        approver=None,
        month=7,
        score=Decimal("90"),
        critical=False,
        manager_comment="Stage 8 manager comment",
    ):
        employee = employee or self.employee
        manager = manager or self.manager
        approver = approver or self.director
        review = create_kpi_review(
            employee=employee,
            period_type=KPIReview.PERIOD_MONTHLY,
            period_start=date(2026, month, 1),
            period_end=date(2026, month, 28),
            review_date=date(2026, month, 15),
            actor=manager,
        )
        entry = review.item_entries.get()
        save_kpi_review_draft(
            review,
            actor=manager,
            entry_values={
                entry.pk: {
                    "actual_value": score,
                    "critical_red": critical,
                    "critical_reason": (
                        "Critical approved quality failure" if critical else ""
                    ),
                    "critical_trigger": "Quality escalation" if critical else "",
                    "item_comment": "Stage 8 approved source",
                }
            },
            manager_comment=manager_comment,
        )
        submit_kpi_review(review, actor=manager)
        start_kpi_review(review, actor=approver)
        return approve_kpi_review(
            review,
            actor=approver,
            comment="Approved Stage 8 result",
        )

    def create_bonus_rule(self):
        profile = KPIBonusWeightProfile.objects.create(
            code="stage8-individual",
            name="Stage 8 Individual Weights",
            version=1,
            team_scope_code="configured-team",
            team_scope_name="Configured Team",
            individual_weight=Decimal("100.00"),
            team_weight=Decimal("0.00"),
            company_weight=Decimal("0.00"),
            effective_start=date(2026, 1, 1),
            created_by=self.admin,
        )
        profile.status = KPIBonusWeightProfile.STATUS_PUBLISHED
        profile.published_by = self.admin
        profile.save()
        rule = KPIBonusRuleSet.objects.create(
            code="stage8-individual",
            name="Stage 8 Individual Bonus",
            version=1,
            weight_profile=profile,
            bonus_enabled=True,
            minimum_score=Decimal("70.00"),
            attendance_multiplier_default=Decimal("1.0000"),
            attendance_multiplier_min=Decimal("0.0000"),
            attendance_multiplier_max=Decimal("1.2000"),
            critical_red_behavior=KPIBonusRuleSet.CRITICAL_BLOCK,
            bonus_floor=Decimal("0.00"),
            bonus_cap=Decimal("1000.00"),
            currency="CAD",
            eligible_employee_statuses=["active"],
            approval_required=True,
            effective_start=date(2026, 1, 1),
            created_by=self.admin,
        )
        rule.status = KPIBonusRuleSet.STATUS_PUBLISHED
        rule.published_by = self.admin
        rule.save()
        return rule

    def test_rule_validation_effective_lookup_and_published_immutability(self):
        self.assertEqual(
            effective_intelligence_rule_set(date(2026, 7, 1)).pk,
            self.intelligence_rules.pk,
        )
        self.intelligence_rules.minimum_score = Decimal("60.00")
        with self.assertRaises(ValidationError):
            self.intelligence_rules.save()

        invalid = KPIIntelligenceRuleSet(
            code="invalid",
            name="Invalid",
            version=1,
            effective_start=date(2026, 1, 1),
            minimum_score=Decimal("70.00"),
            health_green_threshold=Decimal("85.00"),
            critical_red_count=1,
            decline_percentage=Decimal("5.00"),
            trend_periods=1,
            overdue_days=7,
            review_completion_target=Decimal("90.00"),
            improvement_threshold=Decimal("5.00"),
            department_risk_threshold=Decimal("70.00"),
            manager_workload_threshold=5,
            overall_kpi_weight=Decimal("10.00"),
            review_completion_weight=Decimal("20.00"),
            bonus_readiness_weight=Decimal("15.00"),
            improvement_weight=Decimal("10.00"),
            risk_control_weight=Decimal("15.00"),
        )
        with self.assertRaises(ValidationError):
            invalid.full_clean()

    def test_publish_rule_requires_executive_and_records_audit(self):
        draft = KPIIntelligenceRuleSet.objects.create(
            code="next-intelligence",
            name="Next Rules",
            version=1,
            effective_start=date(2027, 1, 1),
            minimum_score=Decimal("70.00"),
            health_green_threshold=Decimal("85.00"),
            critical_red_count=2,
            decline_percentage=Decimal("5.00"),
            trend_periods=3,
            overdue_days=7,
            review_completion_target=Decimal("90.00"),
            improvement_threshold=Decimal("5.00"),
            department_risk_threshold=Decimal("70.00"),
            manager_workload_threshold=5,
            overall_kpi_weight=Decimal("40.00"),
            review_completion_weight=Decimal("20.00"),
            bonus_readiness_weight=Decimal("15.00"),
            improvement_weight=Decimal("10.00"),
            risk_control_weight=Decimal("15.00"),
            created_by=self.admin,
        )
        with self.assertRaises(KPIIntelligencePermissionError):
            publish_intelligence_rule_set(draft, actor=self.manager)
        published = publish_intelligence_rule_set(draft, actor=self.ceo)
        self.assertEqual(published.status, KPIIntelligenceRuleSet.STATUS_PUBLISHED)
        self.assertTrue(
            CRMAuditLog.objects.filter(
                module="kpi_intelligence",
                record_id=str(draft.pk),
                field_name="intelligence_rule_changed",
            ).exists()
        )

    def test_page_requires_authentication_and_loads_role_widgets(self):
        self.assertEqual(
            self.client.get(reverse("kpi_intelligence")).status_code,
            302,
        )
        self.client.force_login(self.employee_user)
        response = self.client.get(reverse("kpi_intelligence"), {"year": 2026})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Executive Intelligence")
        self.assertContains(response, 'data-widget-slug="employee-analytics"')
        self.assertNotContains(response, 'data-widget-slug="critical-alerts"')

    def test_employee_scope_is_own_only(self):
        own = self.approved_period(score=Decimal("91"))
        self.approved_period(
            employee=self.other_employee,
            manager=self.other_manager,
            approver=self.other_director,
            score=Decimal("55"),
        )
        payload = intelligence_widget_payload(
            self.employee_user,
            "employee-analytics",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        self.assertEqual([row["employee_id"] for row in payload["rows"]], [own.employee_id])
        self.assertNotIn(
            "critical-alerts",
            {widget.slug for widget in intelligence_widgets(self.employee_user)},
        )

    def test_manager_director_hr_and_ceo_scopes(self):
        self.approved_period(score=Decimal("92"))
        self.approved_period(
            employee=self.other_employee,
            manager=self.other_manager,
            approver=self.other_director,
            score=Decimal("58"),
        )
        manager = intelligence_widget_payload(
            self.manager,
            "employee-analytics",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        director = intelligence_widget_payload(
            self.director,
            "employee-analytics",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        hr = intelligence_widget_payload(
            self.hr,
            "employee-analytics",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        ceo = intelligence_widget_payload(
            self.ceo,
            "employee-analytics",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        self.assertEqual(manager["total"], 1)
        self.assertEqual(director["total"], 1)
        self.assertEqual(hr["total"], 2)
        self.assertEqual(ceo["total"], 2)

    def test_company_health_uses_configured_rules_and_approved_snapshots(self):
        self.approved_period(month=6, score=Decimal("80"))
        self.approved_period(month=7, score=Decimal("90"))
        payload = intelligence_widget_payload(
            self.ceo,
            "company-health",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        self.assertTrue(payload["snapshot_only"])
        self.assertEqual(payload["rule_version"], 1)
        self.assertEqual(payload["overall_kpi"], "90.00")
        self.assertIn(payload["status"], {"green", "yellow", "red"})

    def test_intelligence_never_calls_stage4_calculation_engine(self):
        self.approved_period(score=Decimal("91"))
        with patch(
            "crm.services.kpi_calculation_engine.calculate_employee_kpi_score",
            side_effect=AssertionError("intelligence recalculated KPI"),
        ):
            payload = intelligence_widget_payload(
                self.ceo,
                "company-health",
                DashboardFilters(year=2026),
                use_cache=False,
            )
        self.assertEqual(payload["overall_kpi"], "91.00")

    def test_critical_red_yellow_and_green_intelligence(self):
        self.approved_period(
            employee=self.other_employee,
            manager=self.other_manager,
            approver=self.other_director,
            month=6,
            score=Decimal("62"),
        )
        critical = self.approved_period(
            employee=self.other_employee,
            manager=self.other_manager,
            approver=self.other_director,
            month=7,
            score=Decimal("95"),
            critical=True,
        )
        self.approved_period(month=6, score=Decimal("95"))
        self.approved_period(month=7, score=Decimal("86"))
        red = intelligence_widget_payload(
            self.ceo,
            "critical-alerts",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        yellow = intelligence_widget_payload(
            self.ceo,
            "yellow-attention",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        green = intelligence_widget_payload(
            self.ceo,
            "green-success",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        self.assertTrue(
            any(
                row["source_record"] == f"kpi_review:{critical.pk}"
                for row in red["rows"]
            )
        )
        self.assertTrue(yellow["rows"])
        self.assertTrue(green["rows"])

    def test_trends_and_insufficient_history(self):
        self.approved_period(month=5, score=Decimal("70"))
        self.approved_period(month=6, score=Decimal("80"))
        self.approved_period(month=7, score=Decimal("90"))
        payload = intelligence_widget_payload(
            self.employee_user,
            "trends",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        self.assertEqual(payload["monthly"]["analysis"]["direction"], "improving")
        self.assertEqual(payload["monthly"]["analysis"]["periods_included"], 3)
        self.assertEqual(
            payload["annual"]["analysis"]["data_quality_warning"],
            "Insufficient History",
        )

    def test_filtered_period_uses_effective_rule_and_history_requirement(self):
        KPIIntelligenceRuleSet.objects.filter(pk=self.intelligence_rules.pk).update(
            effective_end=date(2026, 1, 31),
        )
        payload = intelligence_widget_payload(
            self.ceo,
            "trends",
            DashboardFilters(year=2026, month=1),
            use_cache=False,
        )
        self.assertEqual(payload["rule_version"], self.intelligence_rules.version)
        self.assertEqual(
            payload["monthly"]["analysis"]["data_quality_warning"],
            "Insufficient History",
        )

    def test_location_department_and_manager_analytics(self):
        self.approved_period(score=Decimal("94"))
        self.approved_period(
            employee=self.other_employee,
            manager=self.other_manager,
            approver=self.other_director,
            score=Decimal("74"),
        )
        locations = intelligence_widget_payload(
            self.ceo,
            "location-analytics",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        departments = intelligence_widget_payload(
            self.ceo,
            "department-analytics",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        managers = intelligence_widget_payload(
            self.ceo,
            "manager-analytics",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        location_rows = {row["name"]: row for row in locations["rows"]}
        self.assertEqual(location_rows["Canada"]["score"], "94.00")
        self.assertEqual(location_rows["Bangladesh"]["score"], "74.00")
        self.assertEqual(len(departments["rows"]), 2)
        self.assertEqual(len(managers["rows"]), 2)

    def test_employee_analytics_uses_snapshot_comments_and_items(self):
        self.approved_period(
            score=Decimal("88"),
            manager_comment="<b>Support the next approved target</b>",
        )
        payload = intelligence_widget_payload(
            self.employee_user,
            "employee-analytics",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        row = payload["rows"][0]
        self.assertEqual(row["strongest_kpi"], "Stage 5 Manual KPI")
        self.assertIn(
            "Support the next approved target",
            row["approved_comments"],
        )
        self.assertNotIn("<b>", " ".join(row["approved_comments"]))

    def test_bonus_readiness_uses_immutable_result_and_hides_amount_by_role(self):
        review = self.approved_period(score=Decimal("90"))
        rule = self.create_bonus_rule()
        create_bonus_calculation(
            review,
            rule_set=rule,
            base_bonus_amount=Decimal("500"),
            actor=self.admin,
        )
        manager = intelligence_widget_payload(
            self.manager,
            "bonus-readiness",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        executive = intelligence_widget_payload(
            self.ceo,
            "bonus-readiness",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        self.assertFalse(manager["amounts_visible"])
        self.assertEqual(manager["rows"][0]["amount"], "")
        self.assertTrue(executive["amounts_visible"])
        self.assertEqual(executive["rows"][0]["amount"], "CAD $450.00")

    def test_data_quality_warnings_do_not_convert_missing_to_zero(self):
        self.approved_period(score=Decimal("90"))
        payload = intelligence_widget_payload(
            self.ceo,
            "data-quality",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        codes = {row["code"] for row in payload["warnings"]}
        self.assertIn("missing_review", codes)
        self.assertIn("missing_bonus_rule", codes)
        self.assertIn("insufficient_history", codes)

    def test_action_link_is_signed_scoped_and_audited(self):
        review = self.approved_period(score=Decimal("60"))
        payload = intelligence_widget_payload(
            self.manager,
            "critical-alerts",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        token = payload["rows"][0]["action_link"].rstrip("/").split("/")[-1]
        target = resolve_intelligence_action(self.manager, token)
        self.assertEqual(target, reverse("kpi_review_detail", args=[review.pk]))
        self.assertTrue(
            CRMAuditLog.objects.filter(
                module="kpi_intelligence",
                field_name="critical_action_link_used",
            ).exists()
        )
        with self.assertRaises(KPIIntelligencePermissionError):
            resolve_intelligence_action(self.other_manager, token)

    def test_unauthorized_action_endpoint_is_blocked(self):
        review = self.approved_period(score=Decimal("60"))
        payload = intelligence_widget_payload(
            self.manager,
            "critical-alerts",
            DashboardFilters(year=2026),
            use_cache=False,
        )
        action_link = payload["rows"][0]["action_link"]
        self.client.force_login(self.other_manager)
        response = self.client.get(action_link)
        self.assertEqual(response.status_code, 403)
        self.assertEqual(review.employee_id, self.employee.pk)

    def test_widget_endpoints_enforce_server_permissions_and_audit(self):
        self.approved_period(score=Decimal("60"))
        self.client.force_login(self.employee_user)
        forbidden = self.client.get(
            reverse("kpi_intelligence_widget", args=["critical-alerts"]),
            {"year": 2026},
        )
        self.assertEqual(forbidden.status_code, 403)

        self.client.force_login(self.ceo)
        response = self.client.get(
            reverse("kpi_intelligence_widget", args=["critical-alerts"]),
            {"year": 2026},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["X-KPI-Source"], "approved-snapshots")
        fields = set(
            CRMAuditLog.objects.filter(module="kpi_intelligence").values_list(
                "field_name",
                flat=True,
            )
        )
        self.assertIn("executive_alert_viewed", fields)
        self.assertIn("cached_intelligence_refreshed", fields)

    def test_cache_is_user_scoped_and_cached_widget_uses_zero_queries(self):
        self.approved_period(score=Decimal("90"))
        filters = DashboardFilters(year=2026)
        first = intelligence_widget_payload(
            self.ceo,
            "company-health",
            filters,
        )
        with CaptureQueriesContext(connection) as queries:
            second = intelligence_widget_payload(
                self.ceo,
                "company-health",
                filters,
            )
        self.assertEqual(first["score"], second["score"])
        self.assertTrue(second["_cache_hit"])
        self.assertEqual(len(queries), 0)

    def test_individual_widget_query_budget_and_no_row_driven_growth(self):
        self.approved_period(score=Decimal("90"))
        dashboard_audience(self.ceo)
        with CaptureQueriesContext(connection) as queries:
            payload = intelligence_widget_payload(
                self.ceo,
                "company-health",
                DashboardFilters(year=2026),
                use_cache=False,
            )
        self.assertEqual(payload["overall_kpi"], "90.00")
        self.assertLessEqual(
            len(queries),
            5,
            [query["sql"] for query in queries],
        )

    def test_widget_queries_do_not_grow_with_visible_employee_rows(self):
        self.approved_period(score=Decimal("90"))
        self.approved_period(
            employee=self.other_employee,
            manager=self.other_manager,
            approver=self.other_director,
            score=Decimal("80"),
        )
        dashboard_audience(self.ceo)
        with CaptureQueriesContext(connection) as queries:
            payload = intelligence_widget_payload(
                self.ceo,
                "employee-analytics",
                DashboardFilters(year=2026),
                use_cache=False,
            )
        self.assertEqual(payload["total"], 2)
        self.assertLessEqual(
            len(queries),
            5,
            [query["sql"] for query in queries],
        )

    def test_warm_shell_query_budget_and_mobile_assets(self):
        self.client.force_login(self.employee_user)
        self.client.get(reverse("kpi_intelligence"), {"year": 2026})
        with CaptureQueriesContext(connection) as queries:
            response = self.client.get(
                reverse("kpi_intelligence"),
                {"year": 2026},
            )
        self.assertEqual(response.status_code, 200)
        self.assertLessEqual(
            len(queries),
            10,
            [query["sql"] for query in queries],
        )
        self.assertContains(response, "kpi_intelligence.css")
        self.assertContains(response, "kpi_intelligence.js")
        self.assertContains(response, "data-kpi-intelligence-widget")

    def test_report_permissions_are_role_scoped(self):
        employee_reports = {
            row["code"] for row in allowed_report_types(self.employee_user)
        }
        self.assertEqual(employee_reports, {"employee-performance"})
        with self.assertRaises(KPIIntelligencePermissionError):
            build_kpi_report(
                self.employee_user,
                "executive-summary",
                DashboardFilters(year=2026),
            )

    def test_all_ten_report_types_build_from_authorized_snapshots(self):
        self.approved_period(score=Decimal("90"))
        for report_type in (
            "executive-summary",
            "department-performance",
            "manager-performance",
            "employee-performance",
            "review-completion",
            "critical-red",
            "bonus-readiness",
            "monthly-comparison",
            "quarterly-comparison",
            "annual-comparison",
        ):
            report = build_kpi_report(
                self.ceo,
                report_type,
                DashboardFilters(year=2026),
            )
            self.assertEqual(report.report_type, report_type)
            self.assertIn("digest-verified", report.data_source_note)

    def test_paginated_reports_request_complete_authorized_results(self):
        filters = DashboardFilters(year=2026)
        for report_type, widget_slug in (
            ("employee-performance", "employee-analytics"),
            ("critical-red", "critical-alerts"),
            ("bonus-readiness", "bonus-readiness"),
        ):
            with patch(
                "crm.services.kpi_reporting.intelligence_widget_payload"
            ) as payload:
                payload.return_value = {
                    "kind": "report-test",
                    "rows": [],
                    "score": "",
                    "status": "",
                }
                build_kpi_report(self.ceo, report_type, filters)
                self.assertTrue(
                    any(
                        call.args[1] == widget_slug
                        and call.kwargs.get("page", "missing") is None
                        and call.kwargs.get("use_cache") is False
                        for call in payload.call_args_list
                    ),
                    f"{report_type} did not request the complete authorized set",
                )

    def test_every_executive_intelligence_widget_renders_independently(self):
        self.approved_period(score=Decimal("90"))
        self.client.force_login(self.ceo)
        for widget in intelligence_widgets(self.ceo):
            response = self.client.get(
                reverse("kpi_intelligence_widget", args=[widget.slug]),
                {"year": 2026},
            )
            self.assertEqual(response.status_code, 200, widget.slug)
            self.assertContains(response, widget.title)

    def test_pdf_excel_csv_and_print_exports(self):
        self.approved_period(
            score=Decimal("90"),
            manager_comment="=unsafe formula <script>alert(1)</script>",
        )
        report = build_kpi_report(
            self.ceo,
            "employee-performance",
            DashboardFilters(year=2026),
        )
        pdf = report_pdf_bytes(report)
        xlsx = report_xlsx_bytes(report)
        csv_bytes = report_csv_bytes(report)
        self.assertTrue(pdf.startswith(b"%PDF"))
        workbook = load_workbook(BytesIO(xlsx), read_only=True)
        values = [
            str(cell.value or "")
            for row in workbook.active.iter_rows()
            for cell in row
        ]
        self.assertFalse(any(value.startswith("=") for value in values))
        csv_text = csv_bytes.decode("utf-8-sig")
        self.assertNotIn("<script>", csv_text)
        self.assertNotIn(",=unsafe", csv_text)

        self.client.force_login(self.ceo)
        for export_format, content_type in (
            ("pdf", "application/pdf"),
            (
                "xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            ("csv", "text/csv"),
            ("print", "text/html"),
        ):
            response = self.client.get(
                reverse(
                    "kpi_intelligence_report",
                    args=["employee-performance", export_format],
                ),
                {"year": 2026},
            )
            self.assertEqual(response.status_code, 200, export_format)
            self.assertTrue(
                response.headers["Content-Type"].startswith(content_type)
            )

    def test_export_scope_blocks_guessed_report_and_bonus_amount(self):
        self.approved_period(score=Decimal("90"))
        self.client.force_login(self.employee_user)
        response = self.client.get(
            reverse(
                "kpi_intelligence_report",
                args=["bonus-readiness", "csv"],
            ),
            {"year": 2026},
        )
        self.assertEqual(response.status_code, 403)

        own = self.client.get(
            reverse(
                "kpi_intelligence_report",
                args=["employee-performance", "csv"],
            ),
            {"year": 2026, "employee": self.other_employee.pk},
        )
        self.assertEqual(own.status_code, 200)
        self.assertNotContains(own, self.other_employee.public_name)

    def test_report_and_sensitive_export_audit_events(self):
        review = self.approved_period(score=Decimal("90"))
        create_bonus_calculation(
            review,
            rule_set=self.create_bonus_rule(),
            base_bonus_amount=Decimal("500"),
            actor=self.admin,
        )
        self.client.force_login(self.ceo)
        response = self.client.get(
            reverse(
                "kpi_intelligence_report",
                args=["bonus-readiness", "csv"],
            ),
            {"year": 2026, "employee": self.employee.pk},
        )
        self.assertEqual(response.status_code, 200)
        events = set(
            CRMAuditLog.objects.filter(module="kpi_intelligence").values_list(
                "field_name",
                flat=True,
            )
        )
        self.assertTrue(
            {
                "report_generated",
                "report_exported",
                "sensitive_bonus_report_opened",
                "filtered_employee_report_generated",
            }.issubset(events)
        )

    def test_routes_are_get_only(self):
        self.client.force_login(self.ceo)
        self.assertEqual(
            self.client.post(reverse("kpi_intelligence")).status_code,
            405,
        )
        self.assertEqual(
            self.client.post(
                reverse(
                    "kpi_intelligence_widget",
                    args=["company-health"],
                )
            ).status_code,
            405,
        )
